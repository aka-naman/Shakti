from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
import os
import shutil
from datetime import datetime

OUTPUT_DIR = "generated_notices"

def replace_in_paragraph(para, placeholders):
    for key, value in placeholders.items():
        for run in para.runs:
            if key in run.text:
                run.text = run.text.replace(key, str(value))


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    """Sets compact cell padding (in dxas) for tighter table density."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(margin)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def generate_tbrl_noting(data):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    master_path = os.path.join("masters", "TBRL_Noting_Master.docx")
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    group = data.get('group_name', 'General').replace("/", "-")
    filename = f"TBRL_Noting_{group}_{timestamp}.docx"
    output_path = os.path.join(OUTPUT_DIR, filename)
    
    shutil.copy2(master_path, output_path)
    doc = Document(output_path)
    
    # ── Page Setup Optimization for Single-Page Printing ──
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    
    # Date Calculations
    start_date = datetime.strptime(data['start_date'], '%Y-%m-%d')
    end_date = datetime.strptime(data['end_date'], '%Y-%m-%d')
    days_count = (end_date - start_date).days + 1
    
    # Format course dates in Hindi
    hindi_months = {
        1: "जनवरी", 2: "फ़रवरी", 3: "मार्च", 4: "अप्रैल",
        5: "मई", 6: "जून", 7: "जुलाई", 8: "अगस्त",
        9: "सितंबर", 10: "अक्टूबर", 11: "नवंबर", 12: "दिसंबर"
    }
    start_day = start_date.day
    end_day = end_date.day
    start_month = hindi_months[start_date.month]
    end_month = hindi_months[end_date.month]
    start_year = start_date.year
    end_year = end_date.year
    
    if start_month == end_month and start_year == end_year:
        course_dates_str = f"{start_day}-{end_day} {start_month} {start_year}"
    elif start_year == end_year:
        course_dates_str = f"{start_day} {start_month} से {end_day} {end_month} {start_year}"
    else:
        course_dates_str = f"{start_day} {start_month} {start_year} से {end_day} {end_month} {end_year}"

    # Grammar Engine (Hindi Singular/Plural)
    num_nominees = len(data.get('nominees', []))
    if num_nominees <= 1:
        off_word, ka_ke, hua_hue, nam_word = "अधिकारी", "का", "हुआ", "नामांकन"
    else:
        off_word, ka_ke, hua_hue, nam_word = "अधिकारियों", "के", "हुए", "नामांकनों"
        
    # Placeholders Mapping
    placeholders = {
        "{{REF_NO}}": data.get("ref_no", ""),
        "{{LAB_NAME}}": data.get("lab_name", "टीबीआरएल"),
        "{{CURRENT_YEAR}}": str(datetime.now().year),
        "{{SUBJECT_HINDI}}": data.get("subject_hindi", ""),
        "{{SUBJECT_ENGLISH}}": data.get("subject_english", ""),
        "{{COURSE_TYPE}}": data.get("course_type", ""),
        "{{REFERENCE_TEXT}}": data.get("reference_text", ""),
        "{{REF_DATE}}": data.get("ref_date", ""),
        "{{COURSE_DATES}}": course_dates_str,
        "{{START_DATE}}": start_date.strftime('%d %B %Y'),
        "{{END_DATE}}": end_date.strftime('%d %B %Y'),
        "{{DAYS_COUNT}}": str(days_count),
        "{{ORG_INSTITUTE}}": data.get("org_institute", ""),
        "{{COURSE_TITLE}}": data.get("course_title", ""),
        "{{GROUP_NAME}}": f"{data.get('group_name', '')} Group",
        "{{OFFICIAL_WORD}}": off_word,
        "{{KA_KE}}": ka_ke,
        "{{HUA_HUE}}": hua_hue,
        "{{NAMANKAN_WORD}}": nam_word,
        "{{SIGNATORY_1_NAME}}": data.get("sig1_name", ""),
        "{{SIGNATORY_1_DESIG}}": data.get("sig1_desig", ""),
        "{{SIGNATORY_2_NAME}}": data.get("sig2_name", ""),
        "{{SIGNATORY_2_DESIG}}": data.get("sig2_desig", ""),
    }

    # Process all body paragraphs, optimizing spacing
    table_index = None
    for i, para in enumerate(doc.paragraphs):
        # Apply tight paragraph styling to conserve height
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.line_spacing = 1.05
        
        # Optimize runs font sizes
        for run in para.runs:
            if run.font.size is None or run.font.size > Pt(11):
                run.font.size = Pt(10) # Drop body text to tight 10pt
                
        if "{{COMPLEX_DYNAMIC_TABLE}}" in para.text:
            table_index = i
            para.text = "" 
        else:
            replace_in_paragraph(para, placeholders)

    # 4. Generate Highly Customizable Table
    if table_index is not None and data.get('columns') and data.get('nominees'):
        columns = data['columns']
        nominees = data['nominees']
        
        table = doc.add_table(rows=len(nominees) + 1, cols=len(columns))
        set_table_borders(table, color="A0AEC0", sz="4")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Draw dynamic headers (simple black bold text, no background color)
        hdr_cells = table.rows[0].cells
        for col_idx, col_name in enumerate(columns):
            cell = hdr_cells[col_idx]
            set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            run = p.add_run(col_name)
            run.font.bold = True
            run.font.size = Pt(9)
            
        # Draw dynamic data rows (simple black text, no background color)
        for row_idx, nominee_data in enumerate(nominees, start=1):
            row_cells = table.rows[row_idx].cells
            
            for col_idx, col_name in enumerate(columns):
                cell = row_cells[col_idx]
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                
                # S.N. column center aligned, others left aligned
                if col_name.lower() in ["s.n.", "s.no.", "sl.no.", "cr.sn.", "क्र.सं."]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
                # Auto-shrink font if table has many columns to fit A4 page
                font_size = 8.5 if len(columns) > 5 else 9.5
                
                cell_value = str(nominee_data.get(col_name, ""))
                run = p.add_run(cell_value)
                run.font.size = Pt(font_size)
                
        # Move table to correct position in document XML
        doc.paragraphs[table_index]._element.addnext(table._element)

    doc.save(output_path)
    return output_path, filename


def generate_lecture_noting(data):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    master_path = os.path.join("masters", "Lecture_Noting_Master.docx")
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    group = data.get('group_name', 'General').replace("/", "-")
    filename = f"Lecture_Noting_{group}_{timestamp}.docx"
    output_path = os.path.join(OUTPUT_DIR, filename)
    
    shutil.copy2(master_path, output_path)
    doc = Document(output_path)
    
    # ── Page Setup Optimization for Single-Page Printing ──
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    
    # Date Calculations
    start_date = datetime.strptime(data['start_date'], '%Y-%m-%d')
    end_date = datetime.strptime(data['end_date'], '%Y-%m-%d')
    days_count = (end_date - start_date).days + 1
    
    # Format course dates in Hindi
    hindi_months = {
        1: "जनवरी", 2: "फ़रवरी", 3: "मार्च", 4: "अप्रैल",
        5: "मई", 6: "जून", 7: "जुलाई", 8: "अगस्त",
        9: "सितंबर", 10: "अक्टूबर", 11: "नवंबर", 12: "दिसंबर"
    }
    start_day = start_date.day
    end_day = end_date.day
    start_month = hindi_months[start_date.month]
    end_month = hindi_months[end_date.month]
    start_year = start_date.year
    end_year = end_date.year
    
    if start_month == end_month and start_year == end_year:
        course_dates_str = f"{start_day}-{end_day} {start_month} {start_year}"
    elif start_year == end_year:
        course_dates_str = f"{start_day} {start_month} से {end_day} {end_month} {start_year}"
    else:
        course_dates_str = f"{start_day} {start_month} {start_year} से {end_day} {end_month} {end_year}"

    # Grammar Engine (Hindi Singular/Plural)
    num_nominees = len(data.get('nominees', []))
    if num_nominees <= 1:
        off_word, ka_ke, hua_hue, nam_word = "अधिकारी", "का", "हुआ", "नामांकन"
    else:
        off_word, ka_ke, hua_hue, nam_word = "अधिकारियों", "के", "हुए", "नामांकनों"
        
    # Placeholders Mapping
    placeholders = {
        "{{REF_NO}}": data.get("ref_no", ""),
        "{{LAB_NAME}}": data.get("lab_name", "टीबीआरएल"),
        "{{CURRENT_YEAR}}": str(datetime.now().year),
        "{{SUBJECT_HINDI}}": data.get("subject_hindi", ""),
        "{{SUBJECT_ENGLISH}}": data.get("subject_english", ""),
        "{{COURSE_TYPE}}": data.get("course_type", ""),
        "{{REFERENCE_TEXT}}": data.get("reference_text", ""),
        "{{REF_DATE}}": data.get("ref_date", ""),
        "{{COURSE_DATES}}": course_dates_str,
        "{{START_DATE}}": start_date.strftime('%d %B %Y'),
        "{{END_DATE}}": end_date.strftime('%d %B %Y'),
        "{{DAYS_COUNT}}": str(days_count),
        "{{ORG_INSTITUTE}}": data.get("org_institute", ""),
        "{{COURSE_TITLE}}": data.get("course_title", ""),
        "{{LECTURE_TITLE}}": data.get("lecture_title", ""),
        "{{GROUP_NAME}}": f"{data.get('group_name', '')} Group",
        "{{OFFICIAL_WORD}}": off_word,
        "{{KA_KE}}": ka_ke,
        "{{HUA_HUE}}": hua_hue,
        "{{NAMANKAN_WORD}}": nam_word,
        "{{SIGNATORY_1_NAME}}": data.get("sig1_name", ""),
        "{{SIGNATORY_1_DESIG}}": data.get("sig1_desig", ""),
        "{{SIGNATORY_2_NAME}}": data.get("sig2_name", ""),
        "{{SIGNATORY_2_DESIG}}": data.get("sig2_desig", ""),
    }

    # Process all body paragraphs, optimizing spacing
    table_index = None
    for i, para in enumerate(doc.paragraphs):
        # Apply tight paragraph styling to conserve height
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.line_spacing = 1.05
        
        # Optimize runs font sizes
        for run in para.runs:
            if run.font.size is None or run.font.size > Pt(11):
                run.font.size = Pt(10) # Drop body text to tight 10pt
                
        if "{{COMPLEX_DYNAMIC_TABLE}}" in para.text:
            table_index = i
            para.text = "" 
        else:
            replace_in_paragraph(para, placeholders)

    # Generate Highly Customizable Table
    if table_index is not None and data.get('columns') and data.get('nominees'):
        columns = data['columns']
        nominees = data['nominees']
        
        table = doc.add_table(rows=len(nominees) + 1, cols=len(columns))
        set_table_borders(table, color="A0AEC0", sz="4")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Draw dynamic headers (simple black bold text, no background color)
        hdr_cells = table.rows[0].cells
        for col_idx, col_name in enumerate(columns):
            cell = hdr_cells[col_idx]
            set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            run = p.add_run(col_name)
            run.font.bold = True
            run.font.size = Pt(9)
            
        # Draw dynamic data rows (simple black text, no background color)
        for row_idx, nominee_data in enumerate(nominees, start=1):
            row_cells = table.rows[row_idx].cells
            
            for col_idx, col_name in enumerate(columns):
                cell = row_cells[col_idx]
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                
                # S.N. column center aligned, others left aligned
                if col_name.lower() in ["s.n.", "s.no.", "sl.no.", "cr.sn.", "क्र.सं.", "क्र.स"]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
                # Auto-shrink font if table has many columns to fit A4 page
                font_size = 8.5 if len(columns) > 5 else 9.5
                
                cell_value = str(nominee_data.get(col_name, ""))
                run = p.add_run(cell_value)
                run.font.size = Pt(font_size)
                
        # Move table to correct position in document XML
        doc.paragraphs[table_index]._element.addnext(table._element)

    doc.save(output_path)
    return output_path, filename


def generate_dgmss_noting(data):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    master_path = os.path.join("masters", "DGMSS_Noting_Master.docx")
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    group = data.get('group_name', 'General').replace("/", "-")
    filename = f"DGMSS_Noting_{group}_{timestamp}.docx"
    output_path = os.path.join(OUTPUT_DIR, filename)
    
    shutil.copy2(master_path, output_path)
    doc = Document(output_path)
    
    # ── Page Setup Optimization for Single-Page Printing ──
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    
    # Date Calculations
    start_date = datetime.strptime(data['start_date'], '%Y-%m-%d')
    end_date = datetime.strptime(data['end_date'], '%Y-%m-%d')
    days_count = (end_date - start_date).days + 1
    
    # Format course dates in Hindi
    hindi_months = {
        1: "जनवरी", 2: "फ़रवरी", 3: "मार्च", 4: "अप्रैल",
        5: "मई", 6: "जून", 7: "जुलाई", 8: "अगस्त",
        9: "सितंबर", 10: "अक्टूबर", 11: "नवंबर", 12: "दिसंबर"
    }
    start_day = start_date.day
    end_day = end_date.day
    start_month = hindi_months[start_date.month]
    end_month = hindi_months[end_date.month]
    start_year = start_date.year
    end_year = end_date.year
    
    if start_month == end_month and start_year == end_year:
        course_dates_str = f"{start_day}-{end_day} {start_month} {start_year}"
    elif start_year == end_year:
        course_dates_str = f"{start_day} {start_month} से {end_day} {end_month} {start_year}"
    else:
        course_dates_str = f"{start_day} {start_month} {start_year} से {end_day} {end_month} {end_year}"

    # Grammar Engine (Hindi Singular/Plural)
    num_nominees = len(data.get('nominees', []))
    if num_nominees <= 1:
        off_word, ka_ke, hua_hue, nam_word = "अधिकारी", "का", "हुआ", "नामांकन"
    else:
        off_word, ka_ke, hua_hue, nam_word = "अधिकारियों", "के", "हुए", "नामांकनों"
        
    # Placeholders Mapping
    placeholders = {
        "{{REF_NO}}": data.get("ref_no", ""),
        "{{LAB_NAME}}": data.get("lab_name", "टीबीआरएल"),
        "{{CURRENT_YEAR}}": str(datetime.now().year),
        "{{SUBJECT_HINDI}}": data.get("subject_hindi", ""),
        "{{SUBJECT_ENGLISH}}": data.get("subject_english", ""),
        "{{COURSE_TYPE}}": data.get("course_type", ""),
        "{{REFERENCE_TEXT}}": data.get("reference_text", ""),
        "{{REF_DATE}}": data.get("ref_date", ""),
        "{{COURSE_DATES}}": course_dates_str,
        "{{START_DATE}}": start_date.strftime('%d %B %Y'),
        "{{END_DATE}}": end_date.strftime('%d %B %Y'),
        "{{DAYS_COUNT}}": str(days_count),
        "{{ORG_INSTITUTE}}": data.get("org_institute", ""),
        "{{COURSE_TITLE}}": data.get("course_title", ""),
        "{{GROUP_NAME}}": f"{data.get('group_name', '')} Group",
        "{{OFFICIAL_WORD}}": off_word,
        "{{KA_KE}}": ka_ke,
        "{{HUA_HUE}}": hua_hue,
        "{{NAMANKAN_WORD}}": nam_word,
        "{{SIGNATORY_1_NAME}}": data.get("sig1_name", ""),
        "{{SIGNATORY_1_DESIG}}": data.get("sig1_desig", ""),
        "{{SIGNATORY_2_NAME}}": data.get("sig2_name", ""),
        "{{SIGNATORY_2_DESIG}}": data.get("sig2_desig", ""),
    }

    # Process all body paragraphs, optimizing spacing
    table_index = None
    for i, para in enumerate(doc.paragraphs):
        # Apply tight paragraph styling to conserve height
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.line_spacing = 1.05
        
        # Optimize runs font sizes
        for run in para.runs:
            if run.font.size is None or run.font.size > Pt(11):
                run.font.size = Pt(10) # Drop body text to tight 10pt
                
        if "{{COMPLEX_DYNAMIC_TABLE}}" in para.text:
            table_index = i
            para.text = "" 
        else:
            replace_in_paragraph(para, placeholders)

    # Generate Highly Customizable Table
    if table_index is not None and data.get('columns') and data.get('nominees'):
        columns = data['columns']
        nominees = data['nominees']
        
        table = doc.add_table(rows=len(nominees) + 1, cols=len(columns))
        set_table_borders(table, color="A0AEC0", sz="4")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Draw dynamic headers (simple black bold text, no background color)
        hdr_cells = table.rows[0].cells
        for col_idx, col_name in enumerate(columns):
            cell = hdr_cells[col_idx]
            set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            run = p.add_run(col_name)
            run.font.bold = True
            run.font.size = Pt(9)
            
        # Draw dynamic data rows (simple black text, no background color)
        for row_idx, nominee_data in enumerate(nominees, start=1):
            row_cells = table.rows[row_idx].cells
            
            for col_idx, col_name in enumerate(columns):
                cell = row_cells[col_idx]
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                
                # S.N. column center aligned, others left aligned
                if col_name.lower() in ["s.n.", "s.no.", "sl.no.", "cr.sn.", "क्र.सं.", "क्र.स"]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
                # Auto-shrink font if table has many columns to fit A4 page
                font_size = 8.5 if len(columns) > 5 else 9.5
                
                cell_value = str(nominee_data.get(col_name, ""))
                run = p.add_run(cell_value)
                run.font.size = Pt(font_size)
                
        # Move table to correct position in document XML
        doc.paragraphs[table_index]._element.addnext(table._element)

    doc.save(output_path)
    return output_path, filename


def generate_fee_noting(data):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    master_path = os.path.join("masters", "Fee_Noting_Master.docx")
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    group = data.get('group_name', 'General').replace("/", "-")
    filename = f"Fee_Noting_{group}_{timestamp}.docx"
    output_path = os.path.join(OUTPUT_DIR, filename)
    
    shutil.copy2(master_path, output_path)
    doc = Document(output_path)
    
    # ── Page Setup Optimization for Single-Page Printing ──
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    
    # Date Calculations
    start_date = datetime.strptime(data['start_date'], '%Y-%m-%d')
    end_date = datetime.strptime(data['end_date'], '%Y-%m-%d')
    days_count = (end_date - start_date).days + 1
    
    # Format course dates in Hindi
    hindi_months = {
        1: "जनवरी", 2: "फ़रवरी", 3: "मार्च", 4: "अप्रैल",
        5: "मई", 6: "जून", 7: "जुलाई", 8: "अगस्त",
        9: "सितंबर", 10: "अक्टूबर", 11: "नवंबर", 12: "दिसंबर"
    }
    start_day = start_date.day
    end_day = end_date.day
    start_month = hindi_months[start_date.month]
    end_month = hindi_months[end_date.month]
    start_year = start_date.year
    end_year = end_date.year
    
    if start_month == end_month and start_year == end_year:
        course_dates_str = f"{start_day}-{end_day} {start_month} {start_year}"
    elif start_year == end_year:
        course_dates_str = f"{start_day} {start_month} से {end_day} {end_month} {start_year}"
    else:
        course_dates_str = f"{start_day} {start_month} {start_year} से {end_day} {end_month} {end_year}"

    # Grammar Engine (Hindi Singular/Plural)
    num_nominees = len(data.get('nominees', []))
    if num_nominees <= 1:
        off_word, ka_ke, hua_hue, nam_word = "अधिकारी", "का", "नामांकन", "विवरण"
    else:
        off_word, ka_ke, hua_hue, nam_word = "अधिकारियों", "के", "नामांकनों", "विवरण"
        
    ref_no = data.get("ref_no", "")
    lab_name = data.get("lab_name", "टीबीआरएल")
    if ref_no.startswith(f"{lab_name}/"):
        ref_no = ref_no[len(lab_name)+1:]

    # Placeholders Mapping
    placeholders = {
        "{{REF_NO}}": ref_no,
        "{{LAB_NAME}}": lab_name,
        "{{CURRENT_YEAR}}": str(datetime.now().year),
        "{{SUBJECT_HINDI}}": data.get("subject_hindi", ""),
        "{{SUBJECT_ENGLISH}}": data.get("subject_english", ""),
        "{{COURSE_TYPE}}": data.get("course_type", ""),
        "{{REFERENCE_TEXT}}": data.get("reference_text", ""),
        "{{REF_DATE}}": data.get("ref_date", ""),
        "{{REF_MAIL_DATE}}": data.get("ref_mail_date", ""),
        "{{COURSE_DATES}}": course_dates_str,
        "{{START_DATE}}": start_date.strftime('%d %B %Y'),
        "{{END_DATE}}": end_date.strftime('%d %B %Y'),
        "{{DAYS_COUNT}}": str(days_count),
        "{{ORG_INSTITUTE}}": data.get("org_institute", ""),
        "{{COURSE_TITLE}}": data.get("course_title", ""),
        "{{GROUP_NAME}}": f"{data.get('group_name', '')} Group",
        "{{OFFICIAL_WORD}}": off_word,
        "{{KA_KE}}": ka_ke,
        "{{HUA_HUE}}": hua_hue,
        "{{NAMANKAN_WORD}}": nam_word,
        "{{SIGNATORY_1_NAME}}": data.get("sig1_name", ""),
        "{{SIGNATORY_1_DESIG}}": data.get("sig1_desig", ""),
        "{{SIGNATORY_2_NAME}}": data.get("sig2_name", ""),
        "{{SIGNATORY_2_DESIG}}": data.get("sig2_desig", ""),
    }

    # Process all body paragraphs, optimizing spacing
    table_index = None
    for i, para in enumerate(doc.paragraphs):
        # Apply tight paragraph styling to conserve height
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.line_spacing = 1.05
        
        # Optimize runs font sizes
        for run in para.runs:
            if run.font.size is None or run.font.size > Pt(11):
                run.font.size = Pt(10) # Drop body text to tight 10pt
                
        if "{{COMPLEX_DYNAMIC_TABLE}}" in para.text:
            table_index = i
            para.text = "" 
        else:
            replace_in_paragraph(para, placeholders)

    # Generate Highly Customizable Table
    if table_index is not None and data.get('columns') and data.get('nominees'):
        columns = data['columns']
        nominees = data['nominees']
        
        table = doc.add_table(rows=len(nominees) + 1, cols=len(columns))
        set_table_borders(table, color="A0AEC0", sz="4")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Draw dynamic headers (simple black bold text, no background color)
        hdr_cells = table.rows[0].cells
        for col_idx, col_name in enumerate(columns):
            cell = hdr_cells[col_idx]
            set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            run = p.add_run(col_name)
            run.font.bold = True
            run.font.size = Pt(9)
            
        # Draw dynamic data rows (simple black text, no background color)
        for row_idx, nominee_data in enumerate(nominees, start=1):
            row_cells = table.rows[row_idx].cells
            
            for col_idx, col_name in enumerate(columns):
                cell = row_cells[col_idx]
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                
                # S.N. column center aligned, others left aligned
                if col_name.lower() in ["s.n.", "s.no.", "sl.no.", "cr.sn.", "क्र.सं.", "क्र.स"]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
                # Auto-shrink font if table has many columns to fit A4 page
                font_size = 8.0 if len(columns) > 7 else (8.5 if len(columns) > 5 else 9.5)
                
                cell_value = str(nominee_data.get(col_name, ""))
                run = p.add_run(cell_value)
                run.font.size = Pt(font_size)
                
        # Move table to correct position in document XML
        doc.paragraphs[table_index]._element.addnext(table._element)

    doc.save(output_path)
    return output_path, filename


def generate_cancellation_noting(data):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    master_path = os.path.join("masters", "Cancellation_Noting_Master.docx")
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    group = data.get('group_name', 'General').replace("/", "-")
    filename = f"Cancellation_Noting_{group}_{timestamp}.docx"
    output_path = os.path.join(OUTPUT_DIR, filename)
    
    shutil.copy2(master_path, output_path)
    doc = Document(output_path)
    
    # ── Page Setup Optimization for Single-Page Printing ──
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    
    # Date Calculations
    start_date = datetime.strptime(data['start_date'], '%Y-%m-%d')
    end_date = datetime.strptime(data['end_date'], '%Y-%m-%d')
    days_count = (end_date - start_date).days + 1
    
    # Format course dates in Hindi
    hindi_months = {
        1: "जनवरी", 2: "फ़रवरी", 3: "मार्च", 4: "अप्रैल",
        5: "मई", 6: "जून", 7: "जुलाई", 8: "अगस्त",
        9: "सितंबर", 10: "अक्टूबर", 11: "नवंबर", 12: "दिसंबर"
    }
    start_day = start_date.day
    end_day = end_date.day
    start_month = hindi_months[start_date.month]
    end_month = hindi_months[end_date.month]
    start_year = start_date.year
    end_year = end_date.year
    
    if start_month == end_month and start_year == end_year:
        course_dates_str = f"{start_day}-{end_day} {start_month} {start_year}"
    elif start_year == end_year:
        course_dates_str = f"{start_day} {start_month} से {end_day} {end_month} {start_year}"
    else:
        course_dates_str = f"{start_day} {start_month} {start_year} से {end_day} {end_month} {end_year}"

    # Placeholders Mapping
    ref_no = data.get("ref_no", "")
    lab_name = data.get("lab_name", "टीबीआरएल")
    if ref_no.startswith(f"{lab_name}/"):
        ref_no = ref_no[len(lab_name)+1:]

    placeholders = {
        "{{REF_NO}}": ref_no,
        "{{LAB_NAME}}": lab_name,
        "{{CURRENT_YEAR}}": str(datetime.now().year),
        "{{SUBJECT_HINDI}}": data.get("subject_hindi", ""),
        "{{SUBJECT_ENGLISH}}": data.get("subject_english", ""),
        "{{COURSE_TYPE}}": data.get("course_type", ""),
        "{{REFERENCE_TEXT}}": data.get("reference_text", ""),
        "{{REF_DATE}}": data.get("ref_date", ""),
        "{{REF_MAIL_DATE}}": data.get("ref_mail_date", ""),
        "{{ION_REF_SOURCE}}": data.get("ion_ref_source", ""),
        "{{ION_REF_DATE}}": data.get("ion_ref_date", ""),
        "{{CANCEL_NOMINEE_NAME}}": data.get("cancel_nominee_name", ""),
        "{{CANCEL_GROUP_NAME}}": data.get("cancel_group_name", ""),
        "{{CANCEL_REASON}}": data.get("cancel_reason", ""),
        "{{COURSE_TYPE_SHORT}}": data.get("course_type_short", ""),
        "{{COURSE_DATES}}": course_dates_str,
        "{{START_DATE}}": start_date.strftime('%d %B %Y'),
        "{{END_DATE}}": end_date.strftime('%d %B %Y'),
        "{{DAYS_COUNT}}": str(days_count),
        "{{ORG_INSTITUTE}}": data.get("org_institute", ""),
        "{{COURSE_TITLE}}": data.get("course_title", ""),
        "{{GROUP_NAME}}": f"{data.get('group_name', '')} Group",
        "{{SIGNATORY_1_NAME}}": data.get("sig1_name", ""),
        "{{SIGNATORY_1_DESIG}}": data.get("sig1_desig", ""),
        "{{SIGNATORY_2_NAME}}": data.get("sig2_name", ""),
        "{{SIGNATORY_2_DESIG}}": data.get("sig2_desig", ""),
    }

    # Process all body paragraphs, optimizing spacing
    table_index = None
    for i, para in enumerate(doc.paragraphs):
        # Apply tight paragraph styling to conserve height
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.line_spacing = 1.05
        
        # Optimize runs font sizes
        for run in para.runs:
            if run.font.size is None or run.font.size > Pt(11):
                run.font.size = Pt(10) # Drop body text to tight 10pt
                
        if "{{COMPLEX_DYNAMIC_TABLE}}" in para.text:
            table_index = i
            para.text = "" 
        else:
            replace_in_paragraph(para, placeholders)

    # Generate Highly Customizable Table
    if table_index is not None and data.get('columns') and data.get('nominees'):
        columns = data['columns']
        nominees = data['nominees']
        
        table = doc.add_table(rows=len(nominees) + 1, cols=len(columns))
        set_table_borders(table, color="A0AEC0", sz="4")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Draw dynamic headers (simple black bold text, no background color)
        hdr_cells = table.rows[0].cells
        for col_idx, col_name in enumerate(columns):
            cell = hdr_cells[col_idx]
            set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            run = p.add_run(col_name)
            run.font.bold = True
            run.font.size = Pt(9)
            
        # Draw dynamic data rows (simple black text, no background color)
        for row_idx, nominee_data in enumerate(nominees, start=1):
            row_cells = table.rows[row_idx].cells
            
            for col_idx, col_name in enumerate(columns):
                cell = row_cells[col_idx]
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                
                # S.N. column center aligned, others left aligned
                if col_name.lower() in ["s.n.", "s.no.", "sl.no.", "cr.sn.", "क्र.सं.", "क्र.स"]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
                # Auto-shrink font if table has many columns to fit A4 page
                font_size = 8.0 if len(columns) > 7 else (8.5 if len(columns) > 5 else 9.5)
                
                cell_value = str(nominee_data.get(col_name, ""))
                run = p.add_run(cell_value)
                run.font.size = Pt(font_size)
                
        # Move table to correct position in document XML
        doc.paragraphs[table_index]._element.addnext(table._element)

    doc.save(output_path)
    return output_path, filename